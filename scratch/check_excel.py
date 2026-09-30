import openpyxl

wb = openpyxl.load_workbook(r"C:\Users\Aaryan shukla\Desktop\sales dummy.xlsx", data_only=False)
ws = wb.active

print(f"Sheet Name: {ws.title}")
print(f"Max Row: {ws.max_row}")
print(f"Max Column: {ws.max_column}")

print("\n--- Row 1 (Headers) ---")
print([cell.value for cell in ws[1]])

print("\n--- Row 2 (First Data Row) ---")
print([cell.value for cell in ws[2]])

print("\n--- Row 50 ---")
print([cell.value for cell in ws[50]])

print("\n--- Row 201 (200th Data Row) ---")
print([cell.value for cell in ws[201]])
